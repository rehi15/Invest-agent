#!/usr/bin/env python3
"""
Excel Auto-Modification Agent
Natural language commands → Excel cell/format updates
"""

import argparse
import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import PatternFill
except ImportError:
    print("openpyxl이 설치되지 않았습니다. 설치 중...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.styles import PatternFill

# ── 색상 매핑 ──────────────────────────────────────────────
COLOR_MAP = {
    "노란색": "FFFF00",
    "노랑":   "FFFF00",
    "빨간색": "FF0000",
    "빨강":   "FF0000",
    "파란색": "0000FF",
    "파랑":   "0000FF",
    "초록색": "00FF00",
    "초록":   "00FF00",
    "주황색": "FFA500",
    "주황":   "FFA500",
    "보라색": "800080",
    "보라":   "800080",
    "회색":   "808080",
    "흰색":   "FFFFFF",
    "하늘색": "87CEEB",
}

DEFAULT_EXCEL = Path("sample_data/구매계획_양식.xlsx")


# ── 유틸 ───────────────────────────────────────────────────
def load_workbook(path: Path):
    return openpyxl.load_workbook(path)


def backup(path: Path) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = path.with_name(f"{path.stem}_backup_{ts}{path.suffix}")
    shutil.copy2(path, bak)
    return bak


def col_letter(ws, header_row: int, col_name: str):
    """컬럼 이름 → 열 번호 (1-based)"""
    for cell in ws[header_row]:
        if cell.value and col_name in str(cell.value):
            return cell.column
    return None


def parse_rows(row_str: str):
    """'10행' or '8행~12행' → list of int"""
    m = re.search(r"(\d+)행[~\-～](\d+)행", row_str)
    if m:
        return list(range(int(m.group(1)), int(m.group(2)) + 1))
    m = re.search(r"(\d+)행", row_str)
    if m:
        return [int(m.group(1))]
    return []


def parse_col_letter(col_str: str):
    """'E열' → 'E'"""
    m = re.search(r"([A-Za-z]+)열", col_str)
    return m.group(1).upper() if m else None


# ── 명령 처리 ──────────────────────────────────────────────
def cmd_input_data(ws, cmd: str) -> str:
    """데이터 입력: '구매 계획 시트 10행에 번호=설인-020, 투자규모=50 입력'"""
    rows = parse_rows(cmd)
    if not rows:
        return "❌ 행 번호를 찾지 못했습니다."

    # key=value 쌍 추출
    pairs = re.findall(r"([가-힣A-Za-z0-9_]+)=([^\s,]+)", cmd)
    if not pairs:
        return "❌ '컬럼명=값' 형식을 찾지 못했습니다."

    # 헤더가 있는 행 자동 탐색 (1~7행)
    header_row = None
    for r in range(1, 8):
        for cell in ws[r]:
            if cell.value:
                header_row = r
                break
        if header_row:
            break

    updated = []
    for row in rows:
        for col_name, value in pairs:
            col_idx = col_letter(ws, header_row, col_name) if header_row else None
            if col_idx:
                ws.cell(row=row, column=col_idx, value=value)
                updated.append(f"({row}행, {col_name})={value}")
            else:
                updated.append(f"⚠️ '{col_name}' 컬럼 미발견")

    return "✅ 입력 완료: " + ", ".join(updated)


def cmd_format_color(ws, cmd: str) -> str:
    """배경색 변경: '10행 E열 배경색 노란색' or '10행~12행 배경색 초기화'"""
    rows = parse_rows(cmd)
    if not rows:
        return "❌ 행 번호를 찾지 못했습니다."

    col_ltr = parse_col_letter(cmd)

    if "초기화" in cmd:
        fill = PatternFill(fill_type=None)
        color_name = "초기화"
    else:
        color_name = next((k for k in COLOR_MAP if k in cmd), None)
        if not color_name:
            return f"❌ 지원 색상: {', '.join(COLOR_MAP.keys())}"
        fill = PatternFill(start_color=COLOR_MAP[color_name],
                           end_color=COLOR_MAP[color_name],
                           fill_type="solid")

    updated = []
    for row in rows:
        if col_ltr:
            ws[f"{col_ltr}{row}"].fill = fill
            updated.append(f"{col_ltr}{row}")
        else:
            for cell in ws[row]:
                cell.fill = fill
            updated.append(f"{row}행 전체")

    return f"✅ 배경색 {color_name} 적용: {', '.join(updated)}"


# ── 시트 선택 ──────────────────────────────────────────────
def select_sheet(wb, cmd: str):
    """명령어에서 시트 이름 추론"""
    for name in wb.sheetnames:
        if name in cmd:
            return wb[name]
    # 부분 매칭
    for name in wb.sheetnames:
        if any(token in cmd for token in name.split()):
            return wb[name]
    return wb.active


# ── 단일 명령 실행 ─────────────────────────────────────────
def run_command(cmd: str, excel_path: Path) -> str:
    if not excel_path.exists():
        return f"❌ 파일 없음: {excel_path}"

    wb = load_workbook(excel_path)
    ws = select_sheet(wb, cmd)

    if any(k in cmd for k in ["입력", "수정", "변경"]) and "배경색" not in cmd:
        result = cmd_input_data(ws, cmd)
    elif "배경색" in cmd:
        result = cmd_format_color(ws, cmd)
    else:
        return "❌ 명령을 인식하지 못했습니다. '입력', '수정', '배경색' 키워드를 사용하세요."

    bak = backup(excel_path)
    wb.save(excel_path)
    return f"{result}\n   백업: {bak.name}"


# ── 배치 실행 ──────────────────────────────────────────────
def run_batch(json_path: Path, excel_path: Path):
    cmds = json.loads(json_path.read_text(encoding="utf-8"))
    for i, cmd in enumerate(cmds, 1):
        print(f"[{i}/{len(cmds)}] {cmd}")
        print("   →", run_command(cmd, excel_path))


# ── 대화형 모드 ───────────────────────────────────────────
def interactive_mode(excel_path: Path):
    print(f"Excel 에이전트 시작 (파일: {excel_path})")
    print("종료: exit 또는 quit\n")
    while True:
        try:
            cmd = input("명령> ").strip()
        except (EOFError, KeyboardInterrupt):
            break
        if not cmd:
            continue
        if cmd.lower() in ("exit", "quit"):
            break
        print(run_command(cmd, excel_path))


# ── CLI 진입점 ────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Excel 자동 수정 에이전트")
    parser.add_argument("--file", default=str(DEFAULT_EXCEL), help="대상 Excel 파일")
    parser.add_argument("--cmd", help="단일 자연어 명령")
    parser.add_argument("--batch", help="배치 명령 JSON 파일")
    parser.add_argument("--interactive", action="store_true", help="대화형 모드")
    args = parser.parse_args()

    excel_path = Path(args.file)

    if args.cmd:
        print(run_command(args.cmd, excel_path))
    elif args.batch:
        run_batch(Path(args.batch), excel_path)
    elif args.interactive:
        interactive_mode(excel_path)
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
